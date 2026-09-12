"""
FastAPI Server for AI Personal Tutor (LangGraph).
=================================================
Exposes endpoints to:
1. Start learning sessions with student profile.
2. Inspect live graph state and memory checkpoints.
3. Resume execution from interrupts (Human-in-the-Loop plan approval & quizzes).
4. Serve the modern interactive web frontend.
"""

import os
import sys
import uuid
from typing import Dict, Any, List, Optional
from fastapi import FastAPI, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

# Ensure project root is available
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(CURRENT_DIR, "..", ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from langgraph.types import Command
from ai_tutor.core.graph import tutor_graph, build_ai_tutor_graph
from ai_tutor.core.state import StudentProfile

app = FastAPI(
    title="AI Personal Tutor - LangGraph API",
    description="Interactive Adaptive Tutor powered by LangGraph StateGraph",
    version="1.0.0"
)

# Static files directory
STATIC_DIR = os.path.join(PROJECT_ROOT, "ai_tutor", "static")

@app.middleware("http")
async def add_no_cache_header(request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


# ---------------------------------------------------------------------------
# Request & Response Schemas
# ---------------------------------------------------------------------------

class StartSessionRequest(BaseModel):
    name: str = "Himaja"
    subject: str = "DBMS"
    goal: str = "Prepare for semester exam"
    level: str = "Beginner"
    available_time: Optional[str] = "Self-paced"
    topics: List[str] = Field(default_factory=list)
    max_quiz_attempts: int = 2


class ResumeSessionRequest(BaseModel):
    resume_payload: Dict[str, Any]  # e.g., {"approved": True} or {"answers": [1, 0]}


class SaveNotesRequest(BaseModel):
    notes: str


# Ensure persistent notes directory
NOTES_DIR = os.path.join(PROJECT_ROOT, "ai_tutor", "data", "notes")
os.makedirs(NOTES_DIR, exist_ok=True)


# ---------------------------------------------------------------------------
# Helper to inspect graph state & interrupts
# ---------------------------------------------------------------------------

def extract_state_snapshot(thread_id: str):
    config = {"configurable": {"thread_id": thread_id}}
    state = tutor_graph.get_state(config)
    
    # Check for active interrupt
    current_interrupt = None
    if state.tasks and len(state.tasks) > 0 and len(state.tasks[0].interrupts) > 0:
        current_interrupt = state.tasks[0].interrupts[0].value

    return {
        "thread_id": thread_id,
        "is_finished": len(state.next) == 0 and current_interrupt is None,
        "next_nodes": list(state.next),
        "current_interrupt": current_interrupt,
        "values": state.values
    }


# ---------------------------------------------------------------------------
# API Endpoints
# ---------------------------------------------------------------------------

@app.post("/api/session/start")
def start_session(req: StartSessionRequest):
    """
    Initializes a new tutor session and executes until the first interrupt (Human-in-the-Loop).
    """
    thread_id = str(uuid.uuid4())[:8]
    config = {"configurable": {"thread_id": thread_id}}
    
    initial_input = {
        "student": req.dict(),
        "completed_topics": [],
        "parallel_checks": [],
        "quiz_history": [],
        "max_quiz_attempts": req.max_quiz_attempts,
        "error_count": 0,
        "needs_revision": False,
        "current_difficulty": "Normal" if req.level.lower() == "intermediate" else ("Advanced" if req.level.lower() == "advanced" else "Easy")
    }
    
    # Run graph until it pauses at the first interrupt (Plan Review)
    for _ in tutor_graph.stream(initial_input, config):
        pass
        
    return extract_state_snapshot(thread_id)


@app.post("/api/session/{thread_id}/resume")
def resume_session(thread_id: str, req: ResumeSessionRequest):
    """
    Resumes graph execution from an interrupt (e.g. plan approval or quiz answer submission).
    Continues stepping until the next interrupt or END.
    """
    config = {"configurable": {"thread_id": thread_id}}
    resume_cmd = Command(resume=req.resume_payload)
    
    # Step through until next interrupt or END
    try:
        for _ in tutor_graph.stream(resume_cmd, config):
            pass
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Graph execution error: {str(e)}")
        
    return extract_state_snapshot(thread_id)


@app.get("/api/session/{thread_id}/state")
def get_session_state(thread_id: str):
    """
    Fetches the current state of a student's session from LangGraph memory checkpointer.
    """
    return extract_state_snapshot(thread_id)


@app.get("/api/session/{thread_id}/history")
def get_session_history(thread_id: str):
    """
    Retrieves the checkpoint history for this thread.
    Demonstrates LangGraph Checkpointing & State Persistence!
    """
    config = {"configurable": {"thread_id": thread_id}}
    checkpoints = []
    
    for state_snapshot in tutor_graph.get_state_history(config):
        checkpoints.append({
            "step": state_snapshot.metadata.get("step", 0),
            "next": list(state_snapshot.next),
            "node_history": state_snapshot.values.get("node_history", []),
            "current_topic": state_snapshot.values.get("current_topic"),
            "completed_topics": state_snapshot.values.get("completed_topics", [])
        })
        
    return {"thread_id": thread_id, "checkpoints": checkpoints}


# ---------------------------------------------------------------------------
# Student Study Notes Persistence & Export Endpoints
# ---------------------------------------------------------------------------

@app.post("/api/session/{thread_id}/notes")
def save_session_notes(thread_id: str, req: SaveNotesRequest):
    """
    Persists student study notes to disk for this session.
    """
    notes_file = os.path.join(NOTES_DIR, f"{thread_id}.txt")
    with open(notes_file, "w", encoding="utf-8") as f:
        f.write(req.notes)
        
    latest_file = os.path.join(NOTES_DIR, "latest_notes.txt")
    with open(latest_file, "w", encoding="utf-8") as f:
        f.write(req.notes)
        
    return {"status": "saved", "thread_id": thread_id, "length": len(req.notes)}


@app.post("/api/notes/save")
def save_generic_notes(req: SaveNotesRequest):
    """
    Persists generic notes even before a session is started.
    """
    latest_file = os.path.join(NOTES_DIR, "latest_notes.txt")
    with open(latest_file, "w", encoding="utf-8") as f:
        f.write(req.notes)
    return {"status": "saved", "length": len(req.notes)}


@app.get("/api/session/{thread_id}/notes")
def get_session_notes(thread_id: str):
    """
    Retrieves persisted study notes for this session.
    """
    notes_file = os.path.join(NOTES_DIR, f"{thread_id}.txt")
    if os.path.exists(notes_file):
        with open(notes_file, "r", encoding="utf-8") as f:
            return {"notes": f.read()}
            
    latest_file = os.path.join(NOTES_DIR, "latest_notes.txt")
    if os.path.exists(latest_file):
        with open(latest_file, "r", encoding="utf-8") as f:
            return {"notes": f.read()}
            
    return {"notes": ""}


@app.get("/api/notes/latest")
def get_latest_notes():
    """
    Retrieves the most recently saved notes.
    """
    latest_file = os.path.join(NOTES_DIR, "latest_notes.txt")
    if os.path.exists(latest_file):
        with open(latest_file, "r", encoding="utf-8") as f:
            return {"notes": f.read()}
    return {"notes": ""}


@app.get("/api/session/{thread_id}/notes/download")
def download_session_notes(thread_id: str, format: str = "txt"):
    """
    Exports student notes as a Notepad text file (.txt) or Word document (.docx).
    """
    config = {"configurable": {"thread_id": thread_id}}
    state = tutor_graph.get_state(config)
    values = state.values or {}
    student = values.get("student", {})
    completed = values.get("completed_topics", [])
    
    notes_file = os.path.join(NOTES_DIR, f"{thread_id}.txt")
    notes_text = ""
    if os.path.exists(notes_file):
        with open(notes_file, "r", encoding="utf-8") as f:
            notes_text = f.read()
    else:
        latest_file = os.path.join(NOTES_DIR, "latest_notes.txt")
        if os.path.exists(latest_file):
            with open(latest_file, "r", encoding="utf-8") as f:
                notes_text = f.read()

    safe_subject = student.get("subject", "General_Study").replace(" ", "_").replace("&", "and")
    student_name = student.get("name", "Student")
    subject_title = student.get("subject", "Course")
    goal = student.get("goal", "Study")

    if format.lower() == "docx":
        from docx import Document
        from docx.shared import Pt, RGBColor
        
        doc = Document()
        doc.add_heading("AI Personal Tutor - Study Notes", level=0)
        
        doc.add_heading("Session Details", level=1)
        doc.add_paragraph(f"Student: {student_name}")
        doc.add_paragraph(f"Subject: {subject_title}")
        doc.add_paragraph(f"Goal: {goal}")
        doc.add_paragraph(f"Completed Topics: {', '.join(completed) if completed else 'In Progress'}")
        
        doc.add_heading("My Personal Study Notes", level=1)
        if notes_text.strip():
            doc.add_paragraph(notes_text)
        else:
            doc.add_paragraph("(No custom notes entered)")
            
        doc.add_paragraph("\n---\nExported from SYNXA AI Personal Tutor")
        
        filename = f"Study_Notes_{safe_subject}_{thread_id}.docx"
        export_path = os.path.join(NOTES_DIR, filename)
        doc.save(export_path)
        
        return FileResponse(
            export_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=filename
        )
    else:
        # Default: Notepad Plain Text (.txt)
        filename = f"Study_Notes_{safe_subject}_{thread_id}.txt"
        export_path = os.path.join(NOTES_DIR, filename)
        
        content = f"""=======================================================
AI PERSONAL TUTOR - STUDENT STUDY NOTES (NOTEPAD)
=======================================================
Student:          {student_name}
Subject:          {subject_title}
Goal:             {goal}
Completed Topics: {', '.join(completed) if completed else 'In Progress'}
=======================================================

MY PERSONAL STUDY NOTES:
-------------------------------------------------------
{notes_text if notes_text.strip() else '(No custom notes entered)'}
-------------------------------------------------------

Exported from SYNXA AI Personal Tutor
"""
        with open(export_path, "w", encoding="utf-8") as f:
            f.write(content)
            
        return FileResponse(
            export_path,
            media_type="text/plain",
            filename=filename
        )


@app.get("/api/notes/download")
def download_latest_notes(format: str = "txt"):
    """
    Exports latest notes if no thread is active in Notepad (.txt) or Word (.docx).
    """
    latest_file = os.path.join(NOTES_DIR, "latest_notes.txt")
    notes_text = ""
    if os.path.exists(latest_file):
        with open(latest_file, "r", encoding="utf-8") as f:
            notes_text = f.read()

    if format.lower() == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading("AI Personal Tutor - My Study Notes", level=0)
        if notes_text.strip():
            doc.add_paragraph(notes_text)
        else:
            doc.add_paragraph("(No custom notes entered)")
            
        filename = "My_Study_Notes.docx"
        export_path = os.path.join(NOTES_DIR, filename)
        doc.save(export_path)
        return FileResponse(
            export_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=filename
        )
    else:
        filename = "My_Study_Notes.txt"
        export_path = os.path.join(NOTES_DIR, filename)
        content = f"""=======================================================
AI PERSONAL TUTOR - MY STUDY NOTES (NOTEPAD)
=======================================================

{notes_text if notes_text.strip() else '(No custom notes entered)'}

=======================================================
Exported from SYNXA AI Personal Tutor
"""
        with open(export_path, "w", encoding="utf-8") as f:
            f.write(content)
            
        return FileResponse(
            export_path,
            media_type="text/plain",
            filename=filename
        )


# ---------------------------------------------------------------------------
# Static Web Frontend
# ---------------------------------------------------------------------------

if os.path.exists(STATIC_DIR):
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

@app.get("/")
def serve_index():
    index_path = os.path.join(STATIC_DIR, "index.html")
    if os.path.exists(index_path):
        return FileResponse(
            index_path,
            headers={
                "Cache-Control": "no-cache, no-store, must-revalidate",
                "Pragma": "no-cache",
                "Expires": "0"
            }
        )
    return {"message": "AI Personal Tutor API is running. Build the frontend to view UI."}
